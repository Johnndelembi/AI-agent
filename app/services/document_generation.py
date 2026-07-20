"""Generate downloadable documents from assistant-authored content."""

from __future__ import annotations

import csv
import io
import json
from textwrap import wrap

from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from app.services.file_service import file_service
from app.utils.error_handler import create_400_error


class DocumentGenerationService:
    """Create simple exportable documents for chat users."""

    SUPPORTED_FORMATS = {"txt", "md", "docx", "pdf", "csv", "json"}

    def generate(
        self,
        *,
        user_id: str,
        title: str,
        content: str,
        output_format: str = "docx",
    ) -> dict:
        normalized_format = output_format.lower()
        if normalized_format not in self.SUPPORTED_FORMATS:
            raise create_400_error(
                f"Unsupported output format '{output_format}'. Supported formats: {', '.join(sorted(self.SUPPORTED_FORMATS))}"
            )

        filename = f"{self._slugify(title)}.{normalized_format}"
        bytes_content, mime_type = self._render_document(title, content, normalized_format)
        return file_service.create_generated_file(
            user_id=user_id,
            filename=filename,
            content=bytes_content,
            mime_type=mime_type,
            file_type="generated_document",
            parsed_content=content,
            metadata={
                "title": title,
                "format": normalized_format,
            },
        )

    def _render_document(self, title: str, content: str, output_format: str) -> tuple[bytes, str]:
        if output_format in {"txt", "md"}:
            mime = "text/plain" if output_format == "txt" else "text/markdown"
            return content.encode("utf-8"), mime
        if output_format == "json":
            payload = json.dumps({"title": title, "content": content}, indent=2, ensure_ascii=True)
            return payload.encode("utf-8"), "application/json"
        if output_format == "csv":
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            writer.writerow(["title", "content"])
            writer.writerow([title, content])
            return buffer.getvalue().encode("utf-8"), "text/csv"
        if output_format == "docx":
            from docx import Document

            document = Document()
            document.add_heading(title, level=1)
            for block in self._split_blocks(content):
                document.add_paragraph(block)
            output = io.BytesIO()
            document.save(output)
            return (
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        if output_format == "pdf":
            output = io.BytesIO()
            pdf = canvas.Canvas(output, pagesize=LETTER)
            width, height = LETTER
            y = height - 72
            pdf.setFont("Helvetica-Bold", 16)
            pdf.drawString(72, y, title[:90])
            y -= 28
            pdf.setFont("Helvetica", 11)
            for block in self._split_blocks(content):
                for line in wrap(block, width=95):
                    if y <= 72:
                        pdf.showPage()
                        pdf.setFont("Helvetica", 11)
                        y = height - 72
                    pdf.drawString(72, y, line)
                    y -= 16
                y -= 8
            pdf.save()
            return output.getvalue(), "application/pdf"
        raise create_400_error(f"Unsupported output format '{output_format}'")

    def _split_blocks(self, content: str) -> list[str]:
        blocks = [block.strip() for block in content.split("\n") if block.strip()]
        return blocks or [content.strip()]

    def _slugify(self, title: str) -> str:
        normalized = "".join(ch.lower() if ch.isalnum() else "-" for ch in title)
        parts = [part for part in normalized.split("-") if part]
        return "-".join(parts)[:80] or "document"


document_generation_service = DocumentGenerationService()

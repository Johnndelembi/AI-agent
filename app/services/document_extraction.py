"""Document extraction helpers for uploaded chat files."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from app.config import logger


class DocumentExtractionService:
    """Extract text from common office and plain-text document formats."""

    MAX_EXTRACTED_CHARS = 100_000

    def extract(self, file_path: Path, mime_type: str | None = None) -> str:
        extension = file_path.suffix.lower()

        try:
            if extension in {".txt", ".md", ".markdown", ".json"}:
                return self._extract_plain_text(file_path, extension)
            if extension == ".csv":
                return self._extract_csv(file_path)
            if extension in {".xlsx", ".xls"}:
                return self._extract_spreadsheet(file_path)
            if extension == ".pdf":
                return self._extract_pdf(file_path)
            if extension == ".docx":
                return self._extract_docx(file_path)
            if extension == ".pptx":
                return self._extract_pptx(file_path)
        except Exception as exc:
            logger.warning("Document extraction failed for %s: %s", file_path, exc)
            return ""

        if mime_type and mime_type.startswith("text/"):
            return self._safe_limit(file_path.read_text(encoding="utf-8", errors="ignore"))
        return ""

    def _extract_plain_text(self, file_path: Path, extension: str) -> str:
        raw = file_path.read_text(encoding="utf-8", errors="ignore")
        if extension == ".json":
            try:
                raw = json.dumps(json.loads(raw), indent=2, ensure_ascii=True)
            except json.JSONDecodeError:
                pass
        return self._safe_limit(raw)

    def _extract_csv(self, file_path: Path) -> str:
        with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            rows = list(csv.reader(handle))
        lines = [", ".join(row) for row in rows[:200]]
        return self._safe_limit("\n".join(lines))

    def _extract_spreadsheet(self, file_path: Path) -> str:
        import pandas as pd

        workbook = pd.ExcelFile(file_path)
        sections: list[str] = []
        for sheet_name in workbook.sheet_names[:10]:
            frame = workbook.parse(sheet_name).fillna("")
            preview = frame.head(50).to_csv(index=False).strip()
            sections.append(f"[Sheet: {sheet_name}]\n{preview}")
        return self._safe_limit("\n\n".join(sections))

    def _extract_pdf(self, file_path: Path) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(file_path))
        texts: list[str] = []
        for page in reader.pages[:200]:
            texts.append(page.extract_text() or "")
        return self._safe_limit("\n\n".join(texts))

    def _extract_docx(self, file_path: Path) -> str:
        from docx import Document

        document = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return self._safe_limit("\n".join(paragraphs))

    def _extract_pptx(self, file_path: Path) -> str:
        from pptx import Presentation

        presentation = Presentation(str(file_path))
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides[:100], start=1):
            text_chunks = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text and shape.text.strip()
            ]
            if text_chunks:
                slides.append(f"[Slide {index}]\n" + "\n".join(text_chunks))
        return self._safe_limit("\n\n".join(slides))

    def _safe_limit(self, text: str) -> str:
        normalized = text.strip()
        if len(normalized) <= self.MAX_EXTRACTED_CHARS:
            return normalized
        return normalized[: self.MAX_EXTRACTED_CHARS].rstrip() + "\n...[truncated]"


document_extraction_service = DocumentExtractionService()
